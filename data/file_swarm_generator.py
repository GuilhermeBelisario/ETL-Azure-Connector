from faker import Faker
from openpyxl import Workbook
from dotenv import load_dotenv
from reportlab.lib.pagesizes import letter
from docx import Document
from reportlab.platypus import SimpleDocTemplate, Table
import os
import random

load_dotenv()

fake = Faker()

def generate_random_data_with_errors(num_rows):
    table = []
    for _ in range(num_rows):

        if random.choice([True, False]):
            table.append(fake.word())
            
        else:
            table.append(random.choice([fake.email(), 'email', 'CNPJ']))
    return table

def generate_random_data():

    if random.choice([True, False]):
        return fake.word() 
    
    else:
        return str(random.randint(1000, 9999))

def generate_table(num_rows):

    table = []

    for _ in range(num_rows):

        row = [generate_random_data() for _ in range(3)]  
        table.append(row)

    return table

def generate_pdf_file(table, file_name):
    pdf = SimpleDocTemplate(file_name, pagesize=letter)
    header_pdf_table = [["1", "2", "3"]] + table
    pdf_table = Table(header_pdf_table)
    
    content_pdf = [pdf_table]
    pdf.build(content_pdf)

def create_docx(file_name, info):
    
    doc = Document()  
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = info

    doc.save(file_name)


num_rows = 30
path_dir = os.getenv("BASE_PATH")
output_dir = f"{path_dir}/data/raw"

for i in range(1, 11):
    file_date = fake.date(pattern='%d.%m.%Y')  

    
    table = generate_table(num_rows)
    wb = Workbook()
    ws = wb.active
    ws.append(["1", "2", "3"])
    for row in table:
        ws.append(row)
    file_name_xlsx = f"{output_dir}/{file_date} random{i}.xlsx"
    wb.save(file_name_xlsx)

   
    if i <= 5:
        table_pdf = generate_table(num_rows)
    else:
        table_pdf = generate_random_data_with_errors(num_rows)
    file_name_pdf = f"{output_dir}/{file_date} random{i}.pdf"
    generate_pdf_file(table_pdf, file_name_pdf)

    
    if i <= 5:
        info_docx = generate_random_data()
    else:
        info_docx = generate_random_data_with_errors(num_rows)[0]  
    file_name_docx = f"{output_dir}/{file_date} random{i}.docx"
    create_docx(file_name_docx, info_docx)